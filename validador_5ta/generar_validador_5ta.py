#!/usr/bin/env python3
"""Genera la plantilla Excel "Validador de retenciones de renta de 5ta categoría" (Perú).

Uso:
    python3 generar_validador_5ta.py                      # 100 trabajadores, 1500 filas de ingresos
    python3 generar_validador_5ta.py --trabajadores 300 --filas-ingresos 5000 --sin-ejemplos

Todas las celdas de cálculo se escriben como fórmulas de Excel, de modo que la
plantilla recalcula al modificar los datos de entrada.
"""
import argparse
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------------------
# Datos normativos
# --------------------------------------------------------------------------------------
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Setiembre", "Octubre", "Noviembre", "Diciembre"]
MES_ABR = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Set", "Oct", "Nov", "Dic"]

# Art. 40 del Reglamento de la LIR: divisor y último mes cuyas retenciones se descuentan
DIVISORES = [(12, 0), (12, 0), (12, 0), (9, 3), (8, 4), (8, 4), (8, 4), (5, 7), (4, 8),
             (4, 8), (4, 8), (1, 11)]
REGLA_DIVISOR = [
    "Impuesto anual / 12", "Impuesto anual / 12", "Impuesto anual / 12",
    "(Impuesto anual − retenciones ene–mar) / 9",
    "(Impuesto anual − retenciones ene–abr) / 8",
    "(Impuesto anual − retenciones ene–abr) / 8",
    "(Impuesto anual − retenciones ene–abr) / 8",
    "(Impuesto anual − retenciones ene–jul) / 5",
    "(Impuesto anual − retenciones ene–ago) / 4",
    "(Impuesto anual − retenciones ene–ago) / 4",
    "(Impuesto anual − retenciones ene–ago) / 4",
    "Regularización: impuesto anual − retenciones ene–nov",
]

# Art. 53 de la LIR: escala progresiva acumulativa (límites en UIT)
TRAMOS = [(0, 5, 0.08), (5, 20, 0.14), (20, 35, 0.17), (35, 45, 0.20), (45, None, 0.30)]

UITS = [(2020, 4300, "D.S. N° 380-2019-EF"), (2021, 4400, "D.S. N° 392-2020-EF"),
        (2022, 4600, "D.S. N° 398-2021-EF"), (2023, 4950, "D.S. N° 309-2022-EF"),
        (2024, 5150, "D.S. N° 309-2023-EF"), (2025, 5350, "D.S. N° 260-2024-EF"),
        (2026, 5500, "D.S. N° 301-2025-EF")]

FRECUENCIAS = [
    ("Mensual", 1, "Todos los meses entre Mes inicio y Mes fin"),
    ("Bimestral", 2, "Cada 2 meses desde Mes inicio hasta Mes fin"),
    ("Trimestral", 3, "Cada 3 meses desde Mes inicio hasta Mes fin"),
    ("Cuatrimestral", 4, "Cada 4 meses desde Mes inicio hasta Mes fin"),
    ("Semestral", 6, "Cada 6 meses desde Mes inicio (p. ej. inicio 5 = mayo y noviembre)"),
    ("Julio y Diciembre", 0, "Julio y diciembre (gratificaciones), dentro de Mes inicio–Mes fin"),
    ("Pago único", 12, "Solo en el Mes inicio (pago particular / esporádico)"),
]

TIPOS = [
    ("ORD", "Remuneración ordinaria mensual",
     "Se proyecta: monto del mes × meses que faltan para terminar el ejercicio (incluido el mes)."),
    ("GRA", "Gratificación ordinaria (julio / diciembre)",
     "Se suma a la proyección mientras no haya sido percibida (incluye la del mes de pago)."),
    ("BON", "Bonificación extraordinaria Ley 30334",
     "Igual que GRA; o como ingreso extraordinario si el parámetro de proyección es 'No'."),
    ("EXT", "Ingreso extraordinario / pago particular",
     "No se proyecta: retención adicional en el mes de pago (art. 40 Reglamento LIR)."),
    ("INA", "Inafecto / no constituye renta de 5ta",
     "No integra la renta bruta. Solo informativo."),
    ("RET", "Retención de 5ta efectuada por otro empleador",
     "Se descuenta como retención ya efectuada en el mes registrado."),
]

CONCEPTOS = [
    ("Remuneración básica / sueldo", "ORD", "Sí", "Mensual",
     "Remuneración ordinaria (art. 34 inc. a LIR)."),
    ("Asignación familiar", "ORD", "Sí", "Mensual", "10% de la RMV (Ley N° 25129)."),
    ("Remuneración vacacional", "ORD", "Sí", "Mensual",
     "Sustituye a la remuneración del mes de descanso vacacional."),
    ("Bonificación mensual fija (cargo, convenio, etc.)", "ORD", "Sí", "Mensual",
     "Pago regular de cada mes."),
    ("Alimentación principal en dinero", "ORD", "Sí", "Mensual", "Pago regular de cada mes."),
    ("Horas extras", "ORD", "No", "Pago único",
     "Ordinaria del mes. Marque 'Sí' en computable si es regular para la gratificación."),
    ("Comisiones / destajo", "ORD", "No", "Pago único",
     "Remuneración variable del mes. Registre una fila por mes."),
    ("Gratificación ordinaria (Fiestas Patrias / Navidad)", "GRA", "No", "Julio y Diciembre",
     "Ley N° 27735."),
    ("Bonificación extraordinaria Ley 30334 (9% / 6.75%)", "BON", "No", "Julio y Diciembre",
     "Afecta a renta de 5ta; inafecta a aportaciones (Leyes N° 30334 y 30712)."),
    ("Participación en las utilidades", "EXT", "No", "Pago único", "D. Leg. N° 892."),
    ("Bono / gratificación extraordinaria", "EXT", "No", "Pago único",
     "Pago no regular; procedimiento de ingresos extraordinarios."),
    ("Bono semestral (productividad, metas, etc.)", "EXT", "No", "Semestral",
     "No se proyecta; se grava en cada mes de pago."),
    ("Bono trimestral (ventas, metas, etc.)", "EXT", "No", "Trimestral",
     "No se proyecta; se grava en cada mes de pago."),
    ("Bono por cierre de pliego / convenio colectivo", "EXT", "No", "Pago único", ""),
    ("Reintegro de remuneraciones", "EXT", "No", "Pago único",
     "Montos devengados en meses anteriores pagados en el mes."),
    ("Asignación por escolaridad / bono anual", "EXT", "No", "Pago único", ""),
    ("Vacaciones truncas", "EXT", "No", "Pago único", "Liquidación de beneficios sociales."),
    ("Gratificación trunca", "EXT", "No", "Pago único", "Liquidación de beneficios sociales."),
    ("Remuneración de otro empleador (total del mes)", "ORD", "No", "Mensual",
     "Usar con Origen = 'Otro empleador' (declaración jurada del trabajador)."),
    ("Retención de 5ta efectuada por otro empleador", "RET", "No", "Pago único",
     "Según certificado de rentas y retenciones. Usar con Origen = 'Otro empleador'."),
    ("CTS – Compensación por tiempo de servicios", "INA", "No", "Semestral",
     "Inafecta (art. 18 inc. c LIR)."),
    ("Indemnización por despido arbitrario", "INA", "No", "Pago único",
     "Inafecta dentro del límite legal (art. 18 inc. c LIR)."),
    ("Movilidad supeditada a la asistencia", "INA", "No", "Mensual",
     "Condición de trabajo; no constituye renta de 5ta."),
    ("Viáticos / gastos de representación sustentados", "INA", "No", "Pago único",
     "Gastos para el cumplimiento de funciones; no constituyen renta."),
]

# --------------------------------------------------------------------------------------
# Datos de ejemplo (códigos EJ-xxx). Se pueden omitir con --sin-ejemplos.
# --------------------------------------------------------------------------------------
EJ_TRABAJADORES = [
    dict(cod="EJ-001", dni="40123456", nombre="PÉREZ GARCÍA, ANA (EJEMPLO)",
         cargo="Analista contable", ingreso=date(2019, 3, 1), cese=None, seguro="EsSalud",
         proy_grat="Sí", proy_otro="No", gastos=6000,
         obs="Aumento desde julio, utilidades en abril, bono semestral y trimestral."),
    dict(cod="EJ-002", dni="45678912", nombre="QUISPE MAMANI, LUIS (EJEMPLO)",
         cargo="Asistente de ventas", ingreso=date(2026, 5, 1), cese=None, seguro="EsSalud",
         proy_grat="Sí", proy_otro="No", gastos=3500,
         obs="Ingresó en mayo; trae rentas y retenciones de su empleador anterior (ene–abr)."),
    dict(cod="EJ-003", dni="70112233", nombre="TORRES ROJAS, MARÍA (EJEMPLO)",
         cargo="Jefa de proyectos", ingreso=date(2022, 8, 15), cese=date(2026, 9, 30),
         seguro="EPS", proy_grat="Sí", proy_otro="Sí", gastos=0,
         obs="Trabajo simultáneo en otra empresa; cese el 30/09 con liquidación."),
    dict(cod="EJ-004", dni="09876543", nombre="SALAZAR VEGA, CARLOS (EJEMPLO)",
         cargo="Gerente general", ingreso=date(2015, 1, 5), cese=None, seguro="EsSalud",
         proy_grat="Sí", proy_otro="No", gastos=16500,
         obs="Bono anual por resultados en marzo (alcanza el tramo de 30%)."),
]

# (código, concepto, origen, otro empleador, frecuencia, monto, mes inicio, mes fin)
P, O = "Empresa principal", "Otro empleador"
EJ_INGRESOS = [
    ("EJ-001", "Remuneración básica / sueldo", P, "", "Mensual", 8500, 1, 6),
    ("EJ-001", "Remuneración básica / sueldo", P, "", "Mensual", 9000, 7, 12),
    ("EJ-001", "Asignación familiar", P, "", "Mensual", 113, 1, 12),
    ("EJ-001", "Gratificación ordinaria (Fiestas Patrias / Navidad)", P, "", "Julio y Diciembre", 9113, 1, 12),
    ("EJ-001", "Bonificación extraordinaria Ley 30334 (9% / 6.75%)", P, "", "Julio y Diciembre", 820.17, 1, 12),
    ("EJ-001", "Participación en las utilidades", P, "", "Pago único", 12000, 4, None),
    ("EJ-001", "Bono semestral (productividad, metas, etc.)", P, "", "Semestral", 2500, 6, 12),
    ("EJ-001", "Bono trimestral (ventas, metas, etc.)", P, "", "Trimestral", 1200, 3, 12),
    ("EJ-001", "Horas extras", P, "", "Pago único", 640, 2, None),
    ("EJ-001", "Horas extras", P, "", "Pago único", 415.5, 8, None),
    ("EJ-001", "CTS – Compensación por tiempo de servicios", P, "", "Semestral", 5150, 5, 11),
    ("EJ-001", "Movilidad supeditada a la asistencia", P, "", "Mensual", 300, 1, 12),
    ("EJ-002", "Remuneración de otro empleador (total del mes)", O,
     "20111111111 – COMERCIAL ANTERIOR S.A.", "Mensual", 5000, 1, 4),
    ("EJ-002", "Gratificación trunca", O, "20111111111 – COMERCIAL ANTERIOR S.A.",
     "Pago único", 2725, 4, None),
    ("EJ-002", "Retención de 5ta efectuada por otro empleador", O,
     "20111111111 – COMERCIAL ANTERIOR S.A.", "Pago único", 721.5, 3, None),
    ("EJ-002", "Remuneración básica / sueldo", P, "", "Mensual", 6000, 5, 12),
    ("EJ-002", "Comisiones / destajo", P, "", "Pago único", 850, 6, None),
    ("EJ-002", "Comisiones / destajo", P, "", "Pago único", 1320, 8, None),
    ("EJ-002", "Gratificación ordinaria (Fiestas Patrias / Navidad)", P, "", "Pago único", 2000, 7, None),
    ("EJ-002", "Bonificación extraordinaria Ley 30334 (9% / 6.75%)", P, "", "Pago único", 180, 7, None),
    ("EJ-002", "Gratificación ordinaria (Fiestas Patrias / Navidad)", P, "", "Pago único", 6000, 12, None),
    ("EJ-002", "Bonificación extraordinaria Ley 30334 (9% / 6.75%)", P, "", "Pago único", 540, 12, None),
    ("EJ-003", "Remuneración básica / sueldo", P, "", "Mensual", 15000, 1, 9),
    ("EJ-003", "Gratificación ordinaria (Fiestas Patrias / Navidad)", P, "", "Pago único", 15000, 7, None),
    ("EJ-003", "Bonificación extraordinaria Ley 30334 (9% / 6.75%)", P, "", "Pago único", 1012.5, 7, None),
    ("EJ-003", "Remuneración de otro empleador (total del mes)", O,
     "20222222222 – CONSULTORA ANDINA S.A.C.", "Mensual", 4000, 1, 12),
    ("EJ-003", "Gratificación ordinaria (Fiestas Patrias / Navidad)", O,
     "20222222222 – CONSULTORA ANDINA S.A.C.", "Julio y Diciembre", 4360, 1, 12),
    ("EJ-003", "Gratificación trunca", P, "", "Pago único", 8006.25, 9, None),
    ("EJ-003", "Vacaciones truncas", P, "", "Pago único", 5000, 9, None),
    ("EJ-003", "CTS – Compensación por tiempo de servicios", P, "", "Pago único", 8750, 5, None),
    ("EJ-003", "CTS – Compensación por tiempo de servicios", P, "", "Pago único", 6250, 9, None),
    ("EJ-004", "Remuneración básica / sueldo", P, "", "Mensual", 35000, 1, 12),
    ("EJ-004", "Gratificación ordinaria (Fiestas Patrias / Navidad)", P, "", "Julio y Diciembre", 35000, 1, 12),
    ("EJ-004", "Bonificación extraordinaria Ley 30334 (9% / 6.75%)", P, "", "Julio y Diciembre", 3150, 1, 12),
    ("EJ-004", "Bono / gratificación extraordinaria", P, "", "Pago único", 50000, 3, None),
    ("EJ-004", "Bono por cierre de pliego / convenio colectivo", P, "", "Pago único", 4000, 10, None),
]

# Retenciones de planilla de ejemplo (ene–set; oct–dic aún no pagados).
# Contienen 2 diferencias intencionales: EJ-001 abril (no retuvo sobre las utilidades) y
# EJ-004 marzo (retuvo de más). EJ-003 redondea a soles enteros (dentro de la tolerancia).
EJ_PLANILLA = {
    "EJ-001": [850.93, 933.06, 1026.39, 850.93, 1060.93, 1578.93, 1178.18, 1179.65, 1289.47],
    "EJ-002": [None, None, None, None, 420.10, 524.23, 434.98, 604.90, 420.10],
    "EJ-003": [2531, 2531, 2531, 2532, 2532, 2532, 2641, 2684, 2735],
    "EJ-004": [8488.75, 8488.75, 25000.00, 8320.83, 8320.83, 8320.83, 8320.83, 8320.84,
               8320.84],
}

# --------------------------------------------------------------------------------------
# Estilos
# --------------------------------------------------------------------------------------
FONT = "Arial"
C_AZUL_OSC = "1F3864"
C_HEADER = "1F4E78"
F_BASE = Font(name=FONT, size=10)
F_TITULO = Font(name=FONT, size=14, bold=True, color=C_AZUL_OSC)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_SECCION = Font(name=FONT, size=11, bold=True, color=C_AZUL_OSC)
FILL_HDR = PatternFill("solid", fgColor=C_HEADER)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_CLAVE = PatternFill("solid", fgColor="E2EFDA")
FILL_AUX = PatternFill("solid", fgColor="F2F2F2")
FILL_SECCION = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=False)
AL_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

NF_MONEDA = '#,##0.00;(#,##0.00);"-"'
NF_ENTERO = '0;-0;"-"'
NF_PCT = '0.00%'
NF_PCT0 = '0%'
NF_FECHA = 'dd/mm/yyyy'

GRUPO_FILLS = {
    "id": "404040", "ing": "2E75B6", "proy": "548235", "tramo": "7030A0",
    "ret": "C55A11", "val": "7F6000",
}


def estilo(c, font=F_BASE, fill=None, nf=None, al=None, borde=True):
    c.font = font
    if fill:
        c.fill = fill
    if nf:
        c.number_format = nf
    if al:
        c.alignment = al
    if borde:
        c.border = BORDE


def encabezado(ws, fila, col, texto, fill=FILL_HDR, ancho=None):
    c = ws.cell(row=fila, column=col, value=texto)
    estilo(c, F_HDR, fill, al=AL_C)
    if ancho:
        ws.column_dimensions[get_column_letter(col)].width = ancho
    return c


def titulo(ws, texto, subtitulo=None):
    ws["A1"] = texto
    ws["A1"].font = F_TITULO
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = F_SUB


def nombre(wb, nom, ref):
    wb.defined_names[nom] = DefinedName(nom, attr_text=ref)


def lista(ws, rango, fuente, msg=None):
    dv = DataValidation(type="list", formula1=fuente, allow_blank=True)
    if msg:
        dv.error = msg
        dv.errorTitle = "Valor no válido"
    dv.add(rango)
    ws.add_data_validation(dv)


# --------------------------------------------------------------------------------------
# Generador
# --------------------------------------------------------------------------------------
class Generador:
    def __init__(self, n_trab, n_ing, ejemplos, anio):
        self.N = n_trab
        self.R = n_ing
        self.ejemplos = ejemplos
        self.anio = anio
        self.wb = Workbook()
        self.wb._named_styles["Normal"].font = F_BASE
        # filas de datos
        self.t0, self.t1 = 6, 5 + n_trab              # Trabajadores / Retenciones / Resumen
        self.i0, self.i1 = 6, 5 + n_ing               # Ingresos
        self.c0, self.c1 = 6, 5 + 12 * n_trab         # Cálculo Mensual

    # ---------------------------------------------------------------- Parámetros
    def hoja_parametros(self):
        ws = self.wb.create_sheet("Parámetros")
        titulo(ws, "PARÁMETROS DEL CÁLCULO",
               "Celdas amarillas: valores editables. Tasas, divisores y UIT según normativa vigente.")
        for col, w in zip("ABCDEFGH", [2, 52, 22, 18, 16, 18, 18, 16]):
            ws.column_dimensions[col].width = w

        def seccion(fila, texto):
            ws.cell(row=fila, column=2, value=texto).font = F_SECCION

        seccion(4, "DATOS GENERALES")
        filas = [
            (5, "Empresa (razón social)", "EMPRESA DEMO S.A.C." if self.ejemplos else "", None, None,
             "Dato informativo."),
            (6, "RUC", "20000000001" if self.ejemplos else "", None, None, "Dato informativo."),
            (7, "Ejercicio gravable", self.anio, "0", None, "Año que se valida."),
            (8, "UIT del ejercicio (S/)", '=IFERROR(VLOOKUP(C7,UIT_TABLA,2,FALSE),0)', NF_MONEDA, None,
             "Se obtiene de la tabla 'UIT por ejercicio'. Si sale 0, agregue la UIT del año."),
            (9, "Deducción fija anual (UIT) – art. 46 LIR", 7, "0", "=C9*C8",
             "La aplica el empleador en cada cálculo mensual."),
            (10, "Deducción adicional máxima por gastos (UIT) – art. 46 LIR", 3, "0", "=C10*C8",
             "Solo para estimar la DJ anual (no la aplica el empleador)."),
            (11, "Bonificación extraordinaria Ley 30334 – EsSalud", 0.09, NF_PCT, None,
             "Equivale al aporte a EsSalud sobre la gratificación."),
            (12, "Bonificación extraordinaria Ley 30334 – EPS", 0.0675, NF_PCT, None,
             "Trabajadores afiliados a una EPS."),
            (13, "¿Proyectar la bonificación extraordinaria con las gratificaciones?", "Sí", None, None,
             "Sí: se proyecta como la gratificación. No: se trata como ingreso extraordinario."),
            (14, "Retenciones de meses anteriores que se descuentan", "Planilla (si existe)", None, None,
             "'Planilla (si existe)': usa lo retenido en planilla; si falta, lo calculado."),
            (15, "Tolerancia para validar diferencias (S/)", 1, NF_MONEDA, None,
             "Diferencias menores o iguales se marcan como OK."),
        ]
        ws.cell(row=4, column=4, value="Equivalente S/").font = F_BOLD
        ws.cell(row=4, column=5, value="Comentario").font = F_BOLD
        for fila, etiqueta, valor, nf, equiv, coment in filas:
            estilo(ws.cell(row=fila, column=2, value=etiqueta))
            c = ws.cell(row=fila, column=3, value=valor)
            es_formula = isinstance(valor, str) and valor.startswith("=")
            estilo(c, F_BASE if es_formula else F_INPUT, None if es_formula else FILL_INPUT, nf,
                   Alignment(horizontal="right"))
            if equiv:
                estilo(ws.cell(row=fila, column=4, value=equiv), nf=NF_MONEDA)
            ws.cell(row=fila, column=5, value=coment).font = F_SUB
        nombre(self.wb, "ANIO", "'Parámetros'!$C$7")
        nombre(self.wb, "UIT", "'Parámetros'!$C$8")
        nombre(self.wb, "DED_FIJA", "'Parámetros'!$D$9")
        nombre(self.wb, "DED_ADIC", "'Parámetros'!$D$10")
        nombre(self.wb, "BON_ESSALUD", "'Parámetros'!$C$11")
        nombre(self.wb, "BON_EPS", "'Parámetros'!$C$12")
        nombre(self.wb, "PROY_BON", "'Parámetros'!$C$13")
        nombre(self.wb, "BASE_RET", "'Parámetros'!$C$14")
        nombre(self.wb, "TOL", "'Parámetros'!$C$15")
        dv = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100")
        dv.add("C7")
        ws.add_data_validation(dv)
        lista(ws, "C13", "LST_SINO")
        lista(ws, "C14", "LST_BASE")

        # Escala por tramos
        seccion(17, "ESCALA PROGRESIVA ACUMULATIVA – art. 53 LIR")
        for j, h in enumerate(["Tramo", "Desde (UIT)", "Hasta (UIT)", "Tasa", "Desde (S/)",
                               "Hasta (S/)", "Tasa incremental"]):
            encabezado(ws, 18, 2 + j, h)
        for k, (d, h, t) in enumerate(TRAMOS):
            f = 19 + k
            estilo(ws.cell(row=f, column=2, value=k + 1), al=Alignment(horizontal="center"))
            estilo(ws.cell(row=f, column=3, value=d), F_INPUT, FILL_INPUT, "0")
            estilo(ws.cell(row=f, column=4, value=h), F_INPUT, FILL_INPUT, "0")
            estilo(ws.cell(row=f, column=5, value=t), F_INPUT, FILL_INPUT, NF_PCT0)
            estilo(ws.cell(row=f, column=6, value=f"=C{f}*UIT"), nf=NF_MONEDA)
            estilo(ws.cell(row=f, column=7, value=f'=IF(D{f}="",1E+15,D{f}*UIT)'),
                   nf='[>=1E+14]"En adelante";#,##0.00')
            inc = f"=E{f}" if k == 0 else f"=E{f}-E{f - 1}"
            estilo(ws.cell(row=f, column=8, value=inc), nf=NF_PCT0)
        ws.cell(row=24, column=2, value="Hasta (UIT) en blanco = en adelante. La tasa incremental "
                "se usa para calcular el impuesto acumulado con SUMPRODUCT.").font = F_SUB
        nombre(self.wb, "TR_TASA", "'Parámetros'!$E$19:$E$23")
        nombre(self.wb, "TR_DESDE", "'Parámetros'!$F$19:$F$23")
        nombre(self.wb, "TR_HASTA", "'Parámetros'!$G$19:$G$23")
        nombre(self.wb, "TR_INC", "'Parámetros'!$H$19:$H$23")

        # Divisores
        seccion(26, "DIVISORES MENSUALES – art. 40 Reglamento LIR (D.S. N° 122-94-EF)")
        for j, h in enumerate(["Mes N°", "Mes", "Divisor", "Descuenta retenciones hasta el mes",
                               "Regla"]):
            encabezado(ws, 27, 2 + j, h)
        ws.merge_cells("F27:H27")
        for m in range(12):
            f = 28 + m
            estilo(ws.cell(row=f, column=2, value=m + 1), al=Alignment(horizontal="center"))
            estilo(ws.cell(row=f, column=3, value=MESES[m]))
            estilo(ws.cell(row=f, column=4, value=DIVISORES[m][0]), F_INPUT, FILL_INPUT, "0")
            estilo(ws.cell(row=f, column=5, value=DIVISORES[m][1]), F_INPUT, FILL_INPUT, "0")
            estilo(ws.cell(row=f, column=6, value=REGLA_DIVISOR[m]))
            ws.merge_cells(start_row=f, start_column=6, end_row=f, end_column=8)
        nombre(self.wb, "DIV_N", "'Parámetros'!$D$28:$D$39")
        nombre(self.wb, "DIV_HASTA", "'Parámetros'!$E$28:$E$39")
        nombre(self.wb, "LST_MESES", "'Parámetros'!$C$28:$C$39")

        # UIT
        seccion(41, "UIT POR EJERCICIO")
        for j, h in enumerate(["Año", "UIT (S/)", "Norma"]):
            encabezado(ws, 42, 2 + j, h)
        for k in range(10):
            f = 43 + k
            a, u, n = UITS[k] if k < len(UITS) else (None, None, None)
            estilo(ws.cell(row=f, column=2, value=a), F_INPUT, FILL_INPUT, "0",
                   Alignment(horizontal="center"))
            estilo(ws.cell(row=f, column=3, value=u), F_INPUT, FILL_INPUT, NF_MONEDA)
            estilo(ws.cell(row=f, column=4, value=n), F_INPUT, FILL_INPUT)
        nombre(self.wb, "UIT_TABLA", "'Parámetros'!$B$43:$C$52")

        # Frecuencias
        seccion(54, "FRECUENCIAS DE PAGO (hoja Ingresos)")
        for j, h in enumerate(["Frecuencia", "Cada (meses)", "Meses en que se paga"]):
            encabezado(ws, 55, 2 + j, h)
        ws.merge_cells("D55:H55")
        for k, (fr, cada, desc) in enumerate(FRECUENCIAS):
            f = 56 + k
            estilo(ws.cell(row=f, column=2, value=fr))
            estilo(ws.cell(row=f, column=3, value=cada), al=Alignment(horizontal="center"))
            estilo(ws.cell(row=f, column=4, value=desc))
            ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=8)
        fin = 55 + len(FRECUENCIAS)
        nombre(self.wb, "FREC_TABLA", f"'Parámetros'!$B$56:$C${fin}")
        nombre(self.wb, "LST_FREC", f"'Parámetros'!$B$56:$B${fin}")

        # Tipos
        seccion(64, "TIPOS DE CONCEPTO (hoja Conceptos)")
        for col, h in ((2, "Tipo"), (3, "Descripción"), (5, "Tratamiento en el cálculo")):
            encabezado(ws, 65, col, h)
        estilo(ws.cell(row=65, column=4), F_HDR, FILL_HDR)
        ws.merge_cells("C65:D65")
        ws.merge_cells("E65:H65")
        for k, (t, d, tr) in enumerate(TIPOS):
            f = 66 + k
            estilo(ws.cell(row=f, column=2, value=t), F_BOLD, al=Alignment(horizontal="center",
                                                                          vertical="center"))
            estilo(ws.cell(row=f, column=3, value=d), al=AL_WRAP)
            estilo(ws.cell(row=f, column=5, value=tr), al=AL_WRAP)
            ws.merge_cells(start_row=f, start_column=3, end_row=f, end_column=4)
            ws.merge_cells(start_row=f, start_column=5, end_row=f, end_column=8)
            ws.row_dimensions[f].height = 27
        nombre(self.wb, "LST_TIPO", "'Parámetros'!$B$66:$B$71")

        # Listas
        seccion(73, "LISTAS DESPLEGABLES")
        listas = [("Sí / No", ["Sí", "No"], "LST_SINO"),
                  ("Origen del ingreso", ["Empresa principal", "Otro empleador"], "LST_ORIGEN"),
                  ("Seguro de salud", ["EsSalud", "EPS", "Sin bonificación"], "LST_SEGURO"),
                  ("Retenciones previas", ["Planilla (si existe)", "Calculadas"], "LST_BASE")]
        for j, (h, vals, nom) in enumerate(listas):
            col = 2 + j
            encabezado(ws, 74, col, h)
            for k, v in enumerate(vals):
                estilo(ws.cell(row=75 + k, column=col, value=v))
            L = get_column_letter(col)
            nombre(self.wb, nom, f"'Parámetros'!${L}$75:${L}${74 + len(vals)}")
        ws.freeze_panes = "A4"
        ws.sheet_view.showGridLines = False

    # ---------------------------------------------------------------- Conceptos
    def hoja_conceptos(self):
        ws = self.wb.create_sheet("Conceptos")
        titulo(ws, "CATÁLOGO DE CONCEPTOS",
               "Puede agregar conceptos en las filas libres. El Tipo define el tratamiento en el "
               "cálculo (ver Parámetros > Tipos de concepto).")
        hdr = ["N°", "Concepto", "Tipo", "Computable para gratificación", "Frecuencia sugerida",
               "Comentario / base legal"]
        for j, (h, w) in enumerate(zip(hdr, [5, 52, 8, 16, 18, 70])):
            encabezado(ws, 4, 1 + j, h, ancho=w)
        n_filas = 60
        for k in range(n_filas):
            f = 5 + k
            datos = CONCEPTOS[k] if k < len(CONCEPTOS) else (None,) * 5
            estilo(ws.cell(row=f, column=1, value=k + 1), al=Alignment(horizontal="center"))
            for j, v in enumerate(datos):
                c = ws.cell(row=f, column=2 + j, value=v)
                estilo(c, F_INPUT, FILL_INPUT,
                       al=Alignment(horizontal="center") if j in (1, 2) else None)
        fin = 4 + n_filas
        nombre(self.wb, "CONC_TABLA", f"Conceptos!$B$5:$D${fin}")
        nombre(self.wb, "LST_CONC", f"Conceptos!$B$5:$B${fin}")
        lista(ws, f"C5:C{fin}", "LST_TIPO")
        lista(ws, f"D5:D{fin}", "LST_SINO")
        lista(ws, f"E5:E{fin}", "LST_FREC")
        ws.freeze_panes = "C5"

    # ---------------------------------------------------------------- Trabajadores
    def hoja_trabajadores(self):
        ws = self.wb.create_sheet("Trabajadores")
        titulo(ws, "TRABAJADORES",
               f"Un trabajador por fila (capacidad: {self.N}). Código único. No elimine filas: "
               "borre el contenido de las celdas amarillas.")
        ws["A3"] = ("Las columnas grises se calculan solas. 'Gastos adicionales' (alquiler, "
                    "honorarios, restaurantes, etc.) solo se usan para estimar la DJ anual.")
        ws["A3"].font = F_SUB
        hdr = [
            ("N°", 5), ("Código", 10), ("DNI", 11), ("Apellidos y nombres", 36),
            ("Cargo / área", 20), ("Fecha de ingreso", 12), ("Fecha de cese", 12),
            ("Seguro de salud", 13), ("¿Proyectar gratificaciones?", 13),
            ("¿Proyectar rem. ordinaria de otro empleador?", 15),
            ("Gastos adicionales sustentados – DJ (S/)", 15), ("Observaciones", 40),
            ("Mes inicio cálculo", 9), ("Mes fin cálculo", 9), ("¿Cesa en el ejercicio?", 9),
            ("% bonif. extraord.", 9), ("Primer mes completo", 9), ("Meses comput. 1er sem.", 9),
            ("Meses comput. 2do sem.", 9), ("Validación", 24),
        ]
        for j, (h, w) in enumerate(hdr):
            encabezado(ws, 5, 1 + j, h, ancho=w)
        ws.row_dimensions[5].height = 54
        ejemplos = EJ_TRABAJADORES if self.ejemplos else []
        t0, t1 = self.t0, self.t1
        for k in range(self.N):
            f = t0 + k
            e = ejemplos[k] if k < len(ejemplos) else {}
            estilo(ws.cell(row=f, column=1, value=k + 1), al=Alignment(horizontal="center"))
            vals = [e.get("cod"), e.get("dni"), e.get("nombre"), e.get("cargo"), e.get("ingreso"),
                    e.get("cese"), e.get("seguro"), e.get("proy_grat"), e.get("proy_otro"),
                    e.get("gastos"), e.get("obs")]
            for j, v in enumerate(vals):
                col = 2 + j
                nf = NF_FECHA if col in (6, 7) else (NF_MONEDA if col == 11 else None)
                al = Alignment(horizontal="center") if col in (2, 3, 6, 7, 8, 9, 10) else None
                estilo(ws.cell(row=f, column=col, value=v), F_INPUT, FILL_INPUT, nf, al)
            B = f"$B{f}"
            formulas = {
                13: f'=IF({B}="","",IF($F{f}="",1,IF(YEAR($F{f})<ANIO,1,IF(YEAR($F{f})=ANIO,MONTH($F{f}),13))))',
                14: f'=IF({B}="","",IF($G{f}="",12,IF(YEAR($G{f})>ANIO,12,IF(YEAR($G{f})=ANIO,MONTH($G{f}),0))))',
                15: f'=IF({B}="","",IF($G{f}="","No",IF(YEAR($G{f})=ANIO,"Sí","No")))',
                16: f'=IF({B}="","",IF($H{f}="EPS",BON_EPS,IF($H{f}="Sin bonificación",0,BON_ESSALUD)))',
                17: f'=IF({B}="","",IF($F{f}="",1,IF(YEAR($F{f})<ANIO,1,IF(YEAR($F{f})>ANIO,13,IF(DAY($F{f})=1,MONTH($F{f}),MONTH($F{f})+1)))))',
                18: f'=IF({B}="","",MAX(0,MIN(6,7-$Q{f})))',
                19: f'=IF({B}="","",MAX(0,MIN(6,13-MAX($Q{f},7))))',
                20: (f'=IF({B}="","",IF(COUNTIF($B${t0}:$B${t1},{B})>1,"Código duplicado",'
                     f'IF($D{f}="","Falta nombre",IF(AND($F{f}<>"",$G{f}<>""),'
                     f'IF($G{f}<$F{f},"Cese anterior al ingreso","OK"),"OK"))))'),
            }
            for col, fx in formulas.items():
                nf = NF_PCT if col == 16 else None
                estilo(ws.cell(row=f, column=col, value=fx), fill=FILL_AUX, nf=nf,
                       al=Alignment(horizontal="center"))
        nombre(self.wb, "LST_COD", f"Trabajadores!$B${t0}:$B${t1}")
        lista(ws, f"H{t0}:H{t1}", "LST_SEGURO")
        lista(ws, f"I{t0}:J{t1}", "LST_SINO")
        dv = DataValidation(type="date", operator="between", formula1="DATE(1950,1,1)",
                            formula2="DATE(2100,12,31)", allow_blank=True)
        dv.add(f"F{t0}:G{t1}")
        ws.add_data_validation(dv)
        ws.freeze_panes = "C6"
        ws.auto_filter.ref = f"A5:T{t1}"
        ws.conditional_formatting.add(
            f"T{t0}:T{t1}", FormulaRule(formula=[f'AND($T{t0}<>"",$T{t0}<>"OK")'],
                                        font=Font(name=FONT, color="C00000", bold=True)))
        ws["I5"].comment = Comment("Sí (o vacío): si la gratificación de julio/diciembre aún no se "
                                   "registró en Ingresos, se estima con la remuneración computable "
                                   "del mes × (1 + % bonificación) × meses completos del semestre / 6.",
                                   "Validador 5ta")
        ws["J5"].comment = Comment("Sí: la remuneración ordinaria del otro empleador (empleo "
                                   "simultáneo) se proyecta hasta diciembre. No (o vacío): se suma "
                                   "conforme se percibe.", "Validador 5ta")

    # ---------------------------------------------------------------- Ingresos
    def hoja_ingresos(self):
        ws = self.wb.create_sheet("Ingresos")
        titulo(ws, "REGISTRO DE INGRESOS",
               "Una fila por concepto. Elija la frecuencia y la plantilla distribuye el monto en los "
               "meses de pago (Ene–Dic). Registre el mes en que se paga (criterio de lo percibido).")
        ws["A3"] = ("Ingresos variables (horas extras, comisiones): una fila 'Pago único' por mes. "
                    "Cambio de sueldo: cierre la fila con 'Mes fin' y cree otra con el nuevo monto. "
                    "Otros empleadores: Origen = 'Otro empleador'.")
        ws["A3"].font = F_SUB
        hdr = [("N°", 6), ("Código trabajador", 11), ("Trabajador", 30), ("Concepto", 40),
               ("Tipo", 6), ("Comp. grat.", 7), ("Origen", 17),
               ("Otro empleador (RUC – razón social)", 28), ("Frecuencia", 15),
               ("Monto por pago (S/)", 13), ("Mes inicio / pago (1-12)", 9),
               ("Mes fin (1-12)", 8), ("Clase", 7), ("Cada (meses)", 7)]
        hdr += [(m, 10) for m in MES_ABR] + [("Total año", 12), ("Validación", 34)]
        for j, (h, w) in enumerate(hdr):
            encabezado(ws, 5, 1 + j, h, ancho=w)
        ws.row_dimensions[5].height = 45
        i0, i1 = self.i0, self.i1
        ejemplos = EJ_INGRESOS if self.ejemplos else []
        for k in range(self.R):
            f = i0 + k
            e = ejemplos[k] if k < len(ejemplos) else (None,) * 8
            cod, conc, orig, emp, frec, monto, mi, mf = e
            estilo(ws.cell(row=f, column=1, value=k + 1), al=Alignment(horizontal="center"))
            for col, v, nf in [(2, cod, None), (4, conc, None), (7, orig, None), (8, emp or None, None),
                               (9, frec, None), (10, monto, NF_MONEDA), (11, mi, "0"), (12, mf, "0")]:
                al = Alignment(horizontal="center") if col in (2, 11, 12) else None
                estilo(ws.cell(row=f, column=col, value=v), F_INPUT, FILL_INPUT, nf, al)
            B = f"$B{f}"
            ws.cell(row=f, column=3,
                    value=f'=IF({B}="","",IFERROR(INDEX(Trabajadores!$D${self.t0}:$D${self.t1},'
                          f'MATCH({B},LST_COD,0)),"¿Código?"))')
            ws.cell(row=f, column=5, value=f'=IF($D{f}="","",IFERROR(VLOOKUP($D{f},CONC_TABLA,2,FALSE),""))')
            ws.cell(row=f, column=6, value=f'=IF($D{f}="","",IFERROR(VLOOKUP($D{f},CONC_TABLA,3,FALSE),""))')
            ws.cell(row=f, column=13,
                    value=(f'=IF(OR({B}="",$E{f}=""),"",IF($E{f}="INA","INA",IF($G{f}="Otro empleador",'
                           f'IF($E{f}="RET","RETO",IF($E{f}="ORD","OTO","OTX")),IF($E{f}="RET","ERR",$E{f}))))'))
            ws.cell(row=f, column=14, value=f'=IF($I{f}="","",IFERROR(VLOOKUP($I{f},FREC_TABLA,2,FALSE),""))')
            for m in range(1, 13):
                col = 14 + m
                fx = (f'=IF(OR($M{f}="",$M{f}="ERR",$J{f}="",$K{f}="",$N{f}=""),0,'
                      f'IF($N{f}=0,IF(AND(OR({m}=7,{m}=12),{m}>=$K{f},{m}<=IF($L{f}="",12,$L{f})),$J{f},0),'
                      f'IF(AND({m}>=$K{f},{m}<=IF($L{f}="",12,$L{f}),MOD({m}-$K{f},$N{f})=0),$J{f},0)))')
                estilo(ws.cell(row=f, column=col, value=fx), nf=NF_MONEDA)
            estilo(ws.cell(row=f, column=27, value=f"=SUM(O{f}:Z{f})"), F_BOLD, nf=NF_MONEDA)
            ws.cell(row=f, column=28,
                    value=(f'=IF(AND({B}="",$D{f}="",$J{f}=""),"",IF({B}="","Falta código",'
                           f'IF(ISNA(MATCH({B},LST_COD,0)),"Código no registrado en Trabajadores",'
                           f'IF($E{f}="","Concepto no está en el catálogo",'
                           f'IF($M{f}="ERR","Retención de otro empleador: Origen = Otro empleador",'
                           f'IF(OR($J{f}="",$K{f}=""),"Falta monto o mes",'
                           f'IF(OR($K{f}<1,$K{f}>12),"Mes inicio inválido",'
                           f'IF(AND($L{f}<>"",OR($L{f}<$K{f},$L{f}>12)),"Mes fin inválido",'
                           f'IF($N{f}="","Falta frecuencia",'
                           f'IF(AND($G{f}<>"Otro empleador",SUMPRODUCT(($O{f}:$Z{f}<>0)*'
                           f'(((COLUMN($O{f}:$Z{f})-14)<INDEX(Trabajadores!$M${self.t0}:$M${self.t1},MATCH({B},LST_COD,0)))'
                           f'+((COLUMN($O{f}:$Z{f})-14)>INDEX(Trabajadores!$N${self.t0}:$N${self.t1},MATCH({B},LST_COD,0)))))>0),'
                           f'"Aviso: hay meses fuera del periodo laboral","OK"))))))))))'))
            for col in (3,):
                estilo(ws.cell(row=f, column=col), F_LINK)
            for col in (5, 6, 13, 14):
                estilo(ws.cell(row=f, column=col), fill=FILL_AUX, al=Alignment(horizontal="center"))
            estilo(ws.cell(row=f, column=28))
        nombre(self.wb, "ING_COD", f"Ingresos!$B${i0}:$B${i1}")
        nombre(self.wb, "ING_COMP", f"Ingresos!$F${i0}:$F${i1}")
        nombre(self.wb, "ING_CLASE", f"Ingresos!$M${i0}:$M${i1}")
        lista(ws, f"B{i0}:B{i1}", "LST_COD", "Elija un código registrado en Trabajadores.")
        lista(ws, f"D{i0}:D{i1}", "LST_CONC", "Elija un concepto del catálogo (hoja Conceptos).")
        lista(ws, f"G{i0}:G{i1}", "LST_ORIGEN")
        lista(ws, f"I{i0}:I{i1}", "LST_FREC")
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                            allow_blank=True)
        dv.error = "Ingrese un mes entre 1 y 12."
        dv.add(f"K{i0}:L{i1}")
        ws.add_data_validation(dv)
        ws.freeze_panes = "E6"
        ws.auto_filter.ref = f"A5:AB{i1}"
        ws.conditional_formatting.add(
            f"AB{i0}:AB{i1}", FormulaRule(formula=[f'LEFT($AB{i0},5)="Aviso"'],
                                          font=Font(name=FONT, color="9C5700", bold=True)))
        ws.conditional_formatting.add(
            f"AB{i0}:AB{i1}", FormulaRule(formula=[f'AND($AB{i0}<>"",$AB{i0}<>"OK")'],
                                          font=Font(name=FONT, color="C00000", bold=True)))
        ws.conditional_formatting.add(
            f"O{i0}:Z{i1}", FormulaRule(formula=[f"O{i0}<>0"], fill=PatternFill("solid", fgColor="EAF1FB")))
        ws["G5"].comment = Comment("Vacío o 'Empresa principal' = ingreso pagado por la empresa que "
                                   "retiene. 'Otro empleador' = empleo simultáneo o empleador anterior "
                                   "del mismo ejercicio.", "Validador 5ta")
        ws["M5"].comment = Comment("ORD/GRA/BON/EXT = empresa principal; OTO = otro empleador "
                                   "ordinario; OTX = otro empleador, otros conceptos; RETO = retención "
                                   "de otro empleador; INA = inafecto.", "Validador 5ta")

    # ---------------------------------------------------------------- Retenciones planilla
    def hoja_retenciones(self):
        ws = self.wb.create_sheet("Retenciones Planilla")
        titulo(ws, "RETENCIONES EFECTUADAS EN PLANILLA",
               "Digite la retención de 5ta categoría descontada cada mes (boleta / PLAME). Deje en "
               "blanco los meses sin información.")
        hdr = [("Código", 10), ("Trabajador", 34)] + [(m, 11) for m in MES_ABR] + \
              [("Total", 13), ("Validación", 26)]
        for j, (h, w) in enumerate(hdr):
            encabezado(ws, 5, 1 + j, h, ancho=w)
        t0, t1 = self.t0, self.t1
        ejemplos = list(EJ_PLANILLA.items()) if self.ejemplos else []
        for k in range(self.N):
            f = t0 + k
            cod, vals = ejemplos[k] if k < len(ejemplos) else (None, [])
            estilo(ws.cell(row=f, column=1, value=cod), F_INPUT, FILL_INPUT,
                   al=Alignment(horizontal="center"))
            estilo(ws.cell(row=f, column=2,
                           value=f'=IF($A{f}="","",IFERROR(INDEX(Trabajadores!$D${t0}:$D${t1},'
                                 f'MATCH($A{f},LST_COD,0)),"¿Código?"))'), F_LINK)
            for m in range(12):
                v = vals[m] if m < len(vals) else None
                estilo(ws.cell(row=f, column=3 + m, value=v), F_INPUT, FILL_INPUT, NF_MONEDA)
            estilo(ws.cell(row=f, column=15, value=f"=SUM(C{f}:N{f})"), F_BOLD, nf=NF_MONEDA)
            estilo(ws.cell(row=f, column=16,
                           value=(f'=IF($A{f}="","",IF(ISNA(MATCH($A{f},LST_COD,0)),'
                                  f'"Código no registrado",IF(COUNTIF($A${t0}:$A${t1},$A{f})>1,'
                                  f'"Código duplicado","OK")))')))
        nombre(self.wb, "RP_COD", f"'Retenciones Planilla'!$A${t0}:$A${t1}")
        lista(ws, f"A{t0}:A{t1}", "LST_COD", "Elija un código registrado en Trabajadores.")
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            allow_blank=True)
        dv.add(f"C{t0}:N{t1}")
        ws.add_data_validation(dv)
        ws.freeze_panes = "C6"
        ws.conditional_formatting.add(
            f"P{t0}:P{t1}", FormulaRule(formula=[f'AND($P{t0}<>"",$P{t0}<>"OK")'],
                                        font=Font(name=FONT, color="C00000", bold=True)))

    # ---------------------------------------------------------------- Cálculo mensual
    CALC_COLS = [
        # clave, encabezado, ancho, grupo, formato
        ("n", "N°", 5, "id", "0"),
        ("cod", "Código", 9, "id", None),
        ("nom", "Apellidos y nombres", 28, "id", None),
        ("mes", "Mes N°", 5, "id", "0"),
        ("mesn", "Mes", 5, "id", None),
        ("est", "Situación", 11, "id", None),
        ("ord", "Rem. ordinaria del mes", 12, "ing", NF_MONEDA),
        ("ordgc", "Rem. computable p/ gratificación", 12, "ing", NF_MONEDA),
        ("gra", "Gratificación ordinaria", 12, "ing", NF_MONEDA),
        ("bon", "Bonif. extraord. Ley 30334", 11, "ing", NF_MONEDA),
        ("ext", "Extraordinarios / pagos particulares", 12, "ing", NF_MONEDA),
        ("oto", "Otro empleador: rem. ordinaria", 12, "ing", NF_MONEDA),
        ("otx", "Otro empleador: otros ingresos", 12, "ing", NF_MONEDA),
        ("reto", "Retención de otro empleador", 11, "ing", NF_MONEDA),
        ("ina", "Inafectos (informativo)", 11, "ing", NF_MONEDA),
        ("tot", "Total ingreso afecto del mes", 12, "ing", NF_MONEDA),
        ("mp", "Meses por proyectar (incl. mes)", 9, "proy", NF_ENTERO),
        ("proy", "(a) Rem. ordinaria proyectada", 13, "proy", NF_MONEDA),
        ("grar", "(a) Gratif. y bonif. registradas por percibir", 13, "proy", NF_MONEDA),
        ("grae", "(a) Gratif. y bonif. estimadas por percibir", 13, "proy", NF_MONEDA),
        ("otop", "Otro empleador: rem. ordinaria (proy. o del mes)", 13, "proy", NF_MONEDA),
        ("otxm", "Otro empleador: otros ingresos del mes", 12, "proy", NF_MONEDA),
        ("ant", "(b) Ingresos afectos de meses anteriores", 13, "proy", NF_MONEDA),
        ("rb", "Renta bruta anual proyectada", 14, "proy", NF_MONEDA),
        ("d7", "(−) Deducción 7 UIT", 12, "proy", NF_MONEDA),
        ("rn", "Renta neta anual proyectada", 14, "proy", NF_MONEDA),
        ("t1", "Tramo 1", 11, "tramo", NF_MONEDA),
        ("t2", "Tramo 2", 11, "tramo", NF_MONEDA),
        ("t3", "Tramo 3", 11, "tramo", NF_MONEDA),
        ("t4", "Tramo 4", 11, "tramo", NF_MONEDA),
        ("t5", "Tramo 5", 11, "tramo", NF_MONEDA),
        ("ia", "Impuesto anual proyectado", 13, "tramo", NF_MONEDA),
        ("tm", "Tasa marginal", 8, "tramo", NF_PCT0),
        ("div", "Divisor", 7, "ret", NF_ENTERO),
        ("hasta", "Descuenta ret. hasta mes", 8, "ret", NF_ENTERO),
        ("retp", "(−) Retenciones de meses anteriores", 13, "ret", NF_MONEDA),
        ("reto_ord", "Retención ordinaria del mes", 12, "ret", NF_MONEDA),
        ("extm", "Extraordinarios del mes", 12, "ret", NF_MONEDA),
        ("rnx", "Renta neta con extraordinarios", 14, "ret", NF_MONEDA),
        ("iax", "Impuesto con extraordinarios", 13, "ret", NF_MONEDA),
        ("retx", "Retención adicional por extraordinarios", 12, "ret", NF_MONEDA),
        ("retc", "RETENCIÓN CALCULADA DEL MES", 13, "ret", NF_MONEDA),
        ("dev", "Exceso a devolver (dic. / cese)", 12, "ret", NF_MONEDA),
        ("retpl", "Retención según planilla", 12, "val", NF_MONEDA),
        ("dif", "Diferencia (planilla − calculada)", 12, "val", NF_MONEDA),
        ("res", "Resultado de la validación", 17, "val", None),
        ("rcons", "Retención considerada (empresa)", 12, "val", NF_MONEDA),
        ("rct", "Retención considerada total (incl. otro empleador)", 13, "val", NF_MONEDA),
    ]
    GRUPOS = {"id": "IDENTIFICACIÓN", "ing": "INGRESOS DEL MES (desde hoja Ingresos)",
              "proy": "PROYECCIÓN ANUAL – art. 40 inc. a) y b)",
              "tramo": "IMPUESTO ANUAL PROYECTADO POR TRAMO – art. 53 LIR",
              "ret": "RETENCIÓN DEL MES – divisores y extraordinarios",
              "val": "VALIDACIÓN CONTRA PLANILLA"}

    def col(self, clave):
        return self.cc[clave]

    def hoja_calculo(self):
        ws = self.wb.create_sheet("Cálculo Mensual")
        self.cc = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(self.CALC_COLS)}
        C = self.cc
        titulo(ws, "CÁLCULO MENSUAL DE RETENCIONES DE 5TA CATEGORÍA",
               "12 filas por trabajador (una por mes). Hoja de fórmulas: no editar. Use el filtro "
               "de 'Código' o la hoja 'Ficha Trabajador' para ver a una persona.")
        # fila de grupos
        grupos = []
        for i, (k, h, w, g, nf) in enumerate(self.CALC_COLS):
            if not grupos or grupos[-1][0] != g:
                grupos.append([g, i + 1, i + 1])
            else:
                grupos[-1][2] = i + 1
        for g, a, b in grupos:
            fill = PatternFill("solid", fgColor=GRUPO_FILLS[g])
            encabezado(ws, 4, a, self.GRUPOS[g], fill)
            for cc in range(a + 1, b + 1):
                estilo(ws.cell(row=4, column=cc), F_HDR, fill)
            if b > a:
                ws.merge_cells(start_row=4, start_column=a, end_row=4, end_column=b)
        for i, (k, h, w, g, nf) in enumerate(self.CALC_COLS):
            fill = PatternFill("solid", fgColor=GRUPO_FILLS[g])
            if k.startswith("t") and k[1:].isdigit():
                t = int(k[1:])
                h = (f'="Tramo {t}: "&TEXT(\'Parámetros\'!$E${18 + t},"0%")')
            encabezado(ws, 5, i + 1, h, fill, w)
        ws.row_dimensions[5].height = 66
        ws.row_dimensions[4].height = 18

        rp_cols = {m: get_column_letter(2 + m) for m in range(1, 13)}  # Retenciones Planilla C..N
        ing_cols = {m: get_column_letter(14 + m) for m in range(1, 13)}  # Ingresos O..Z
        i0, i1 = self.i0, self.i1
        t0, t1 = self.t0, self.t1

        for n in range(1, self.N + 1):
            r0 = self.c0 + (n - 1) * 12
            r1 = r0 + 11
            tr = t0 + n - 1

            def T(colT):
                return f"Trabajadores!${colT}${tr}"

            for m in range(1, 13):
                r = r0 + m - 1
                B = f"$B{r}"
                F = f"${C['est']}{r}"
                act = f'OR({F}="Activo",{F}="Cese")'
                final = "TRUE" if m == 12 else f'{F}="Cese"'
                v = lambda k: f"{C[k]}{r}"  # noqa: E731
                ingc = ing_cols[m]
                rng = f"Ingresos!${ingc}${i0}:${ingc}${i1}"

                def sumifs(clase, extra=""):
                    return (f'=IF({B}="",0,SUMIFS({rng},ING_COD,{B},ING_CLASE,"{clase}"{extra}))')

                bon_pct = f'(1+IF(PROY_BON="Sí",{T("P")},0))'
                grae_terms = []
                if m < 7:
                    grae_terms.append(f'IF({C["gra"]}{r0 + 6}=0,{v("ordgc")}*{bon_pct}*{T("R")}/6*({T("N")}>=7),0)')
                if m < 12:
                    grae_terms.append(f'IF({C["gra"]}{r1}=0,{v("ordgc")}*{bon_pct}*{T("S")}/6*({T("N")}>=12),0)')
                grae = ("=0" if not grae_terms else
                        f'=IF(OR({F}<>"Activo",{T("I")}="No"),0,{"+".join(grae_terms)})')
                mes_rng = f"${C['mes']}{r}:${C['mes']}{r1}"
                if m == 12:
                    dev = f'=IF({act},MAX(0,{v("retp")}-{v("iax")}),0)'
                else:
                    dev = f'=IF({F}="Cese",MAX(0,{v("retp")}-{v("iax")}),0)'
                rp_rng = f"'Retenciones Planilla'!${rp_cols[m]}${t0}:${rp_cols[m]}${t1}"
                f = {
                    "n": n,
                    "cod": f'=IF(Trabajadores!$B${tr}="","",Trabajadores!$B${tr})',
                    "nom": f'=IF({B}="","",Trabajadores!$D${tr})',
                    "mes": m,
                    "mesn": MES_ABR[m - 1],
                    "est": (f'=IF({B}="","",IF(OR({m}<{T("M")},{m}>{T("N")}),"No vinculado",'
                            f'IF(AND({m}={T("N")},{T("O")}="Sí"),"Cese","Activo")))'),
                    "ord": sumifs("ORD"),
                    "ordgc": sumifs("ORD", ',ING_COMP,"Sí"'),
                    "gra": sumifs("GRA"),
                    "bon": sumifs("BON"),
                    "ext": sumifs("EXT"),
                    "oto": sumifs("OTO"),
                    "otx": sumifs("OTX"),
                    "reto": sumifs("RETO"),
                    "ina": sumifs("INA"),
                    "tot": f'={v("ord")}+{v("gra")}+{v("bon")}+{v("ext")}+{v("oto")}+{v("otx")}',
                    "mp": f'=IF({F}="Activo",{13 - m},IF({F}="Cese",1,0))',
                    "proy": f'={v("ord")}*{v("mp")}',
                    "grar": (f'=IF({F}="Activo",SUMIFS({C["gra"]}{r}:{C["gra"]}{r1},{mes_rng},"<="&{T("N")})'
                             f'+IF(PROY_BON="Sí",SUMIFS({C["bon"]}{r}:{C["bon"]}{r1},{mes_rng},"<="&{T("N")}),0),'
                             f'IF({F}="Cese",{v("gra")}+IF(PROY_BON="Sí",{v("bon")},0),0))'),
                    "grae": grae,
                    "otop": (f'=IF({F}="Activo",IF({T("J")}="Sí",{v("oto")}*{13 - m},{v("oto")}),'
                             f'IF({F}="Cese",{v("oto")},0))'),
                    "otxm": f'=IF({act},{v("otx")},0)',
                    "ant": ("=0" if m == 1 else
                            f'=IF({act},SUM({C["tot"]}{r0}:{C["tot"]}{r - 1}),0)'),
                    "rb": (f'=IF({act},{v("proy")}+{v("grar")}+{v("grae")}+{v("otop")}+{v("otxm")}'
                           f'+{v("ant")},0)'),
                    "d7": f'=IF({act},DED_FIJA,0)',
                    "rn": f'=MAX(0,{v("rb")}-{v("d7")})',
                    "ia": f'=SUM({C["t1"]}{r}:{C["t5"]}{r})',
                    "tm": f'=IF({v("rnx")}<=0,"-",INDEX(TR_TASA,MATCH({v("rnx")},TR_DESDE,1)))',
                    "div": f'=IF({F}="Cese",1,IF({F}="Activo",INDEX(DIV_N,{m}),0))',
                    "hasta": f'=IF({F}="Cese",{m - 1},IF({F}="Activo",INDEX(DIV_HASTA,{m}),0))',
                    "retp": ("=0" if m == 1 else
                             f'=IF({act},SUMIFS({C["rct"]}{r0}:{C["rct"]}{r - 1},'
                             f'{C["mes"]}{r0}:{C["mes"]}{r - 1},"<="&{v("hasta")}),0)'),
                    "reto_ord": f'=IF({act},MAX(0,({v("ia")}-{v("retp")})/{v("div")}),0)',
                    "extm": f'=IF({act},{v("ext")}+IF(PROY_BON="Sí",0,{v("bon")}),0)',
                    "rnx": f'=IF({v("extm")}>0,MAX(0,{v("rb")}+{v("extm")}-{v("d7")}),{v("rn")})',
                    "iax": f'=SUMPRODUCT(({v("rnx")}>TR_DESDE)*({v("rnx")}-TR_DESDE)*TR_INC)',
                    "retx": (f'=IF({v("extm")}>0,MAX(0,{v("iax")}-IF({final},MAX({v("ia")},{v("retp")}),'
                             f'{v("ia")})),0)'),
                    "retc": f'=IF({act},ROUND({v("reto_ord")}+{v("retx")},2),0)',
                    "dev": dev,
                    "retpl": (f'=IF({B}="","",IF(COUNTIFS(RP_COD,{B},{rp_rng},"<>")=0,"",'
                              f'SUMIFS({rp_rng},RP_COD,{B})))'),
                    "dif": f'=IF({v("retpl")}="","",{v("retpl")}-{v("retc")})',
                    "res": (f'=IF({B}="","",IF(NOT({act}),IF(N({v("retpl")})=0,"—","Revisar: sin vínculo"),'
                            f'IF({v("retpl")}="",IF({v("retc")}=0,"No corresponde","Sin dato planilla"),IF(ABS({v("dif")})<=TOL,"OK",'
                            f'IF({v("dif")}>0,"Retuvo de más","Retuvo de menos")))))'),
                    "rcons": (f'=IF(AND(BASE_RET="Planilla (si existe)",{v("retpl")}<>""),'
                              f'{v("retpl")},{v("retc")})'),
                    "rct": f'={v("rcons")}+{v("reto")}',
                }
                for t in range(1, 6):
                    p = 18 + t
                    f[f"t{t}"] = (f"=IF({v('rn')}=0,0,MAX(0,MIN({v('rn')},'Parámetros'!$G${p})"
                                  f"-'Parámetros'!$F${p})*'Parámetros'!$E${p})")
                for i, (k, h, w, g, nf) in enumerate(self.CALC_COLS):
                    c = ws.cell(row=r, column=i + 1, value=f[k])
                    font = F_BASE
                    fill = None
                    if k in ("cod", "nom"):
                        font = F_LINK
                    elif k == "retc":
                        font, fill = F_BOLD, FILL_CLAVE
                    elif k in ("ia", "rb", "rn"):
                        font = F_BOLD
                    al = Alignment(horizontal="center") if k in ("n", "mes", "mesn", "est", "mp", "div",
                                                                   "hasta", "tm", "res") else None
                    estilo(c, font, fill, nf, al)
        last = get_column_letter(len(self.CALC_COLS))
        rng_all = f"A{self.c0}:{last}{self.c1}"
        ws.conditional_formatting.add(
            rng_all, FormulaRule(formula=[f'${C["est"]}{self.c0}="No vinculado"'],
                                 font=Font(name=FONT, color="A6A6A6")))
        ws.conditional_formatting.add(
            rng_all, FormulaRule(formula=[f"MOD($A{self.c0},2)=0"],
                                 fill=PatternFill("solid", fgColor="F7F9FC")))
        res_rng = f"{C['res']}{self.c0}:{C['res']}{self.c1}"
        r0 = f"${C['res']}{self.c0}"
        for texto, color, fondo in [("OK", "006100", "C6EFCE"), ("No corresponde", "006100", "C6EFCE"),
                                    ("Retuvo de menos", "9C0006", "FFC7CE"),
                                    ("Retuvo de más", "9C5700", "FFEB9C"),
                                    ("Revisar: sin vínculo", "9C0006", "FFC7CE")]:
            ws.conditional_formatting.add(
                res_rng, FormulaRule(formula=[f'{r0}="{texto}"'], font=Font(name=FONT, color=color, bold=True),
                                     fill=PatternFill("solid", fgColor=fondo)))
        ws.freeze_panes = f"{C['ord']}{self.c0}"
        ws.auto_filter.ref = f"A5:{last}{self.c1}"
        ws.column_dimensions.group(C["ord"], C["tot"], outline_level=1, hidden=False)
        ws.column_dimensions.group(C["t1"], C["t5"], outline_level=1, hidden=False)
        ws.sheet_properties.outlinePr.summaryRight = False

    # ---------------------------------------------------------------- Resumen anual
    RES_COLS = [
        ("n", "N°", 5, None), ("cod", "Código", 9, None), ("nom", "Apellidos y nombres", 30, None),
        ("per", "Periodo con vínculo", 12, None),
        ("ord", "Rem. ordinaria", 13, NF_MONEDA), ("grat", "Gratificaciones y bonif. Ley 30334", 15, NF_MONEDA),
        ("ext", "Extraordinarios / pagos particulares", 15, NF_MONEDA),
        ("otr", "Ingresos de otros empleadores", 13, NF_MONEDA),
        ("rb", "RENTA BRUTA ANUAL", 14, NF_MONEDA), ("ina", "Inafectos (informativo)", 12, NF_MONEDA),
        ("d7", "(−) Deducción 7 UIT", 12, NF_MONEDA), ("rn", "RENTA NETA IMPONIBLE", 14, NF_MONEDA),
        ("t1", "Tramo 1", 11, NF_MONEDA), ("t2", "Tramo 2", 11, NF_MONEDA), ("t3", "Tramo 3", 11, NF_MONEDA),
        ("t4", "Tramo 4", 11, NF_MONEDA), ("t5", "Tramo 5", 11, NF_MONEDA),
        ("ia", "IMPUESTO ANUAL", 13, NF_MONEDA), ("te", "Tasa efectiva", 9, NF_PCT),
        ("tm", "Tasa marginal", 8, NF_PCT0),
        ("retc", "Retenciones calculadas (empresa)", 13, NF_MONEDA),
        ("reto", "Retenciones de otros empleadores", 12, NF_MONEDA),
        ("retpl", "Retenciones según planilla (empresa)", 13, NF_MONEDA),
        ("dif", "Diferencia planilla − calculada (meses con dato)", 12, NF_MONEDA),
        ("rcons", "Retenciones consideradas (total)", 13, NF_MONEDA),
        ("saldo", "Impuesto anual − retenciones consideradas", 13, NF_MONEDA),
        ("obs", "Meses observados", 9, "0"),
        ("gastos", "Gastos adicionales informados", 12, NF_MONEDA),
        ("d3", "(−) Deducción adicional (máx. 3 UIT)", 12, NF_MONEDA),
        ("rndj", "Renta neta DJ anual", 13, NF_MONEDA), ("iadj", "Impuesto anual DJ", 12, NF_MONEDA),
        ("saldodj", "Saldo DJ: a pagar (+) / a favor (−)", 13, NF_MONEDA),
        ("estado", "Estado", 25, None),
    ]

    def hoja_resumen(self):
        ws = self.wb.create_sheet("Resumen Anual")
        R = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(self.RES_COLS)}
        self.rc = R
        C = self.cc
        titulo(ws, "RESUMEN ANUAL POR TRABAJADOR Y POR TRAMOS",
               "Renta e impuesto anual con los ingresos registrados (sin proyección), retenciones "
               "calculadas vs. planilla y estimación de la DJ anual con la deducción adicional de hasta 3 UIT.")
        grupos = [("IDENTIFICACIÓN", "n", "per", "id"), ("RENTA ANUAL", "ord", "rn", "proy"),
                  ("IMPUESTO ANUAL POR TRAMO – art. 53 LIR", "t1", "tm", "tramo"),
                  ("RETENCIONES Y VALIDACIÓN", "retc", "obs", "ret"),
                  ("ESTIMACIÓN DJ ANUAL (deducción adicional hasta 3 UIT)", "gastos", "estado", "val")]
        keys = [k for k, *_r in self.RES_COLS]
        for texto, a, b, g in grupos:
            ia, ib = keys.index(a) + 1, keys.index(b) + 1
            fill = PatternFill("solid", fgColor=GRUPO_FILLS[g])
            encabezado(ws, 4, ia, texto, fill)
            for cc in range(ia + 1, ib + 1):
                estilo(ws.cell(row=4, column=cc), F_HDR, fill)
            ws.merge_cells(start_row=4, start_column=ia, end_row=4, end_column=ib)
        for i, (k, h, w, nf) in enumerate(self.RES_COLS):
            if k.startswith("t") and k[1:].isdigit():
                t = int(k[1:])
                h = f'="Tramo {t}: "&TEXT(\'Parámetros\'!$E${18 + t},"0%")'
            encabezado(ws, 5, i + 1, h, ancho=w)
        ws.row_dimensions[5].height = 66
        t0, t1 = self.t0, self.t1
        CM = "'Cálculo Mensual'!"
        for n in range(1, self.N + 1):
            f = t0 + n - 1
            r0 = self.c0 + (n - 1) * 12
            r1 = r0 + 11
            B = f"$B{f}"

            def blk(k):
                return f"{CM}${C[k]}${r0}:${C[k]}${r1}"

            def S(k):
                return f"SUM({blk(k)})"

            v = lambda k: f"{R[k]}{f}"  # noqa: E731
            fx = {
                "n": n,
                "cod": f'=IF(Trabajadores!$B${f}="","",Trabajadores!$B${f})',
                "nom": f'=IF({B}="","",Trabajadores!$D${f})',
                "per": (f'=IF({B}="","",IF(Trabajadores!$N${f}<Trabajadores!$M${f},"Sin vínculo",'
                        f'INDEX({{"Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Set","Oct","Nov","Dic"}},'
                        f'Trabajadores!$M${f})&" – "&INDEX({{"Ene","Feb","Mar","Abr","May","Jun","Jul","Ago",'
                        f'"Set","Oct","Nov","Dic"}},Trabajadores!$N${f})))'),
                "ord": f"={S('ord')}",
                "grat": f"={S('gra')}+{S('bon')}",
                "ext": f"={S('ext')}",
                "otr": f"={S('oto')}+{S('otx')}",
                "rb": f"={v('ord')}+{v('grat')}+{v('ext')}+{v('otr')}",
                "ina": f"={S('ina')}",
                "d7": f'=IF({B}="",0,DED_FIJA)',
                "rn": f"=MAX(0,{v('rb')}-{v('d7')})",
                "ia": f"=SUM({R['t1']}{f}:{R['t5']}{f})",
                "te": f'=IF(OR({B}="",{v("rb")}=0),"",{v("ia")}/{v("rb")})',
                "tm": f'=IF({B}="","",IF({v("rn")}<=0,"-",INDEX(TR_TASA,MATCH({v("rn")},TR_DESDE,1))))',
                "retc": f"={S('retc')}",
                "reto": f"={S('reto')}",
                "retpl": f"={S('retpl')}",
                "dif": f'=IF(COUNT({blk("retpl")})=0,"",SUM({blk("dif")}))',
                "rcons": f"={S('rct')}",
                "saldo": f"={v('ia')}-{v('rcons')}",
                "obs": f'=IF({B}="","",COUNTIF({blk("res")},"Retuvo*")+COUNTIF({blk("res")},"Revisar*"))',
                "gastos": f'=IF({B}="",0,N(Trabajadores!$K${f}))',
                "d3": f"=MIN({v('gastos')},DED_ADIC,MAX(0,{v('rb')}-{v('d7')}))",
                "rndj": f"=MAX(0,{v('rb')}-{v('d7')}-{v('d3')})",
                "iadj": f"=SUMPRODUCT(({v('rndj')}>TR_DESDE)*({v('rndj')}-TR_DESDE)*TR_INC)",
                "saldodj": f"={v('iadj')}-{v('rcons')}",
                "estado": (f'=IF({B}="","",IF(N({v("obs")})>0,"Revisar ("&{v("obs")}&" mes/es)",'
                           f'IF(COUNTIF({blk("res")},"Sin dato planilla")>0,"OK (hay meses sin dato)","OK")))'),
            }
            for t in range(1, 6):
                p = 18 + t
                fx[f"t{t}"] = (f"=IF({v('rn')}=0,0,MAX(0,MIN({v('rn')},'Parámetros'!$G${p})"
                               f"-'Parámetros'!$F${p})*'Parámetros'!$E${p})")
            for i, (k, h, w, nf) in enumerate(self.RES_COLS):
                c = ws.cell(row=f, column=i + 1, value=fx[k])
                font = F_LINK if k in ("cod", "nom") else (F_BOLD if k in ("rb", "rn", "ia") else F_BASE)
                fill = FILL_CLAVE if k in ("ia", "saldodj") else None
                al = Alignment(horizontal="center") if k in ("n", "per", "tm", "obs", "estado") else None
                estilo(c, font, fill, nf, al)
        # totales
        ft = t1 + 1
        estilo(ws.cell(row=ft, column=3, value="TOTALES"), F_BOLD, FILL_SECCION)
        for i, (k, h, w, nf) in enumerate(self.RES_COLS):
            col = get_column_letter(i + 1)
            if nf == NF_MONEDA:
                estilo(ws.cell(row=ft, column=i + 1, value=f"=SUM({col}{t0}:{col}{t1})"), F_BOLD,
                       FILL_SECCION, nf)
            elif i + 1 != 3:
                estilo(ws.cell(row=ft, column=i + 1), F_BOLD, FILL_SECCION)
        est = f"{R['estado']}{t0}:{R['estado']}{t1}"
        ws.conditional_formatting.add(est, FormulaRule(formula=[f'LEFT(${R["estado"]}{t0},7)="Revisar"'],
                                                       font=Font(name=FONT, color="9C0006", bold=True),
                                                       fill=PatternFill("solid", fgColor="FFC7CE")))
        ws.conditional_formatting.add(est, FormulaRule(formula=[f'LEFT(${R["estado"]}{t0},2)="OK"'],
                                                       font=Font(name=FONT, color="006100", bold=True),
                                                       fill=PatternFill("solid", fgColor="C6EFCE")))
        ws.freeze_panes = "D6"
        ws.auto_filter.ref = f"A5:{get_column_letter(len(self.RES_COLS))}{t1}"

    # ---------------------------------------------------------------- Ficha
    def hoja_ficha(self):
        ws = self.wb.create_sheet("Ficha Trabajador")
        C, R = self.cc, self.rc
        titulo(ws, "FICHA DE CÁLCULO DE RETENCIONES – 5TA CATEGORÍA",
               "Seleccione el código del trabajador en la celda amarilla. Muestra el cálculo de los "
               "12 meses y el impuesto anual por tramos.")
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 50
        for j in range(3, 16):
            ws.column_dimensions[get_column_letter(j)].width = 13.5
        ws["B4"] = "Código del trabajador:"
        ws["B4"].font = F_BOLD
        cod_ini = EJ_TRABAJADORES[0]["cod"] if self.ejemplos else None
        estilo(ws["C4"], Font(name=FONT, size=11, bold=True, color="0000FF"), FILL_INPUT,
               al=Alignment(horizontal="center"))
        ws["C4"] = cod_ini
        lista(ws, "C4", "LST_COD", "Elija un código registrado en Trabajadores.")
        ws["E4"] = "N° interno:"
        ws["E4"].font = F_SUB
        ws["F4"] = f'=IF($C$4="","",IFERROR(MATCH($C$4,LST_COD,0),""))'
        ws["G4"] = "Fila inicial en Cálculo Mensual:"
        ws["G4"].font = F_SUB
        ws["J4"] = f'=IF($F$4="","",{self.c0}+($F$4-1)*12)'
        for ref in ("F4", "J4"):
            estilo(ws[ref], fill=FILL_AUX, al=Alignment(horizontal="center"))
        T = "Trabajadores!"
        datos = [
            ("Apellidos y nombres", f'=IF($F$4="","",INDEX({T}$D${self.t0}:$D${self.t1},$F$4))', None),
            ("DNI", f'=IF($F$4="","",INDEX({T}$C${self.t0}:$C${self.t1},$F$4)&"")', None),
            ("Fecha de ingreso", f'=IF($F$4="","",INDEX({T}$F${self.t0}:$F${self.t1},$F$4))', NF_FECHA),
            ("Fecha de cese", f'=IF($F$4="","",IF(INDEX({T}$G${self.t0}:$G${self.t1},$F$4)="","—",'
                              f'INDEX({T}$G${self.t0}:$G${self.t1},$F$4)))', NF_FECHA),
            ("Seguro de salud / % bonif. extraordinaria",
             f'=IF($F$4="","",INDEX({T}$H${self.t0}:$H${self.t1},$F$4)&" / "&'
             f'TEXT(INDEX({T}$P${self.t0}:$P${self.t1},$F$4),"0.00%"))', None),
            ("Ejercicio / UIT", '=ANIO&" / S/ "&TEXT(UIT,"#,##0")', None),
        ]
        for k, (et, fx, nf) in enumerate(datos):
            f = 6 + k
            estilo(ws.cell(row=f, column=2, value=et), F_BOLD, FILL_SECCION)
            c = ws.cell(row=f, column=3, value=fx)
            estilo(c, nf=nf, al=Alignment(horizontal="left"))
            ws.merge_cells(start_row=f, start_column=3, end_row=f, end_column=7)

        # Tabla mensual
        fh = 13
        encabezado(ws, fh, 2, "Concepto")
        for m in range(12):
            encabezado(ws, fh, 3 + m, MES_ABR[m])
        encabezado(ws, fh, 15, "Total año")
        filas = [
            ("sec", "INGRESOS DEL MES"),
            ("est", "Situación", None, False),
            ("ord", "Remuneración ordinaria", NF_MONEDA, True),
            ("gra", "Gratificación ordinaria", NF_MONEDA, True),
            ("bon", "Bonificación extraordinaria Ley 30334", NF_MONEDA, True),
            ("ext", "Extraordinarios / pagos particulares", NF_MONEDA, True),
            ("oto", "Otro empleador: remuneración ordinaria", NF_MONEDA, True),
            ("otx", "Otro empleador: otros ingresos", NF_MONEDA, True),
            ("tot", "Total ingreso afecto del mes", NF_MONEDA, True),
            ("ina", "Inafectos (informativo)", NF_MONEDA, True),
            ("sec", "PROYECCIÓN ANUAL (art. 40 inc. a y b Reglamento LIR)"),
            ("mp", "Meses por proyectar (incluye el mes)", NF_ENTERO, False),
            ("proy", "(a) Remuneración ordinaria proyectada", NF_MONEDA, False),
            ("grar", "(a) Gratificaciones y bonif. registradas por percibir", NF_MONEDA, False),
            ("grae", "(a) Gratificaciones y bonif. estimadas por percibir", NF_MONEDA, False),
            ("otop", "Otro empleador: rem. ordinaria (proyectada o del mes)", NF_MONEDA, False),
            ("otxm", "Otro empleador: otros ingresos del mes", NF_MONEDA, False),
            ("ant", "(b) Ingresos afectos de meses anteriores", NF_MONEDA, False),
            ("rb", "Renta bruta anual proyectada", NF_MONEDA, False),
            ("d7", "(−) Deducción 7 UIT", NF_MONEDA, False),
            ("rn", "Renta neta anual proyectada", NF_MONEDA, False),
            ("sec", "IMPUESTO ANUAL PROYECTADO POR TRAMO (art. 53 LIR)"),
            ("t1", "Tramo 1", NF_MONEDA, False),
            ("t2", "Tramo 2", NF_MONEDA, False),
            ("t3", "Tramo 3", NF_MONEDA, False),
            ("t4", "Tramo 4", NF_MONEDA, False),
            ("t5", "Tramo 5", NF_MONEDA, False),
            ("ia", "Impuesto anual proyectado", NF_MONEDA, False),
            ("tm", "Tasa marginal", NF_PCT0, False),
            ("sec", "RETENCIÓN DEL MES"),
            ("div", "Divisor", NF_ENTERO, False),
            ("retp", "(−) Retenciones de meses anteriores", NF_MONEDA, False),
            ("reto_ord", "Retención ordinaria del mes", NF_MONEDA, True),
            ("extm", "Extraordinarios del mes", NF_MONEDA, True),
            ("retx", "Retención adicional por extraordinarios", NF_MONEDA, True),
            ("retc", "RETENCIÓN CALCULADA DEL MES", NF_MONEDA, True),
            ("dev", "Exceso a devolver (dic. / cese)", NF_MONEDA, True),
            ("reto", "Retención efectuada por otro empleador", NF_MONEDA, True),
            ("sec", "VALIDACIÓN CONTRA PLANILLA"),
            ("retpl", "Retención según planilla", NF_MONEDA, True),
            ("dif", "Diferencia (planilla − calculada)", NF_MONEDA, True),
            ("res", "Resultado", None, False),
        ]
        f = fh + 1
        CM = "'Cálculo Mensual'!"
        self.ficha_filas = {}
        for item in filas:
            if item[0] == "sec":
                c = ws.cell(row=f, column=2, value=item[1])
                estilo(c, F_SECCION, FILL_SECCION)
                for j in range(3, 16):
                    estilo(ws.cell(row=f, column=j), fill=FILL_SECCION)
                f += 1
                continue
            k, et, nf, total = item
            if k.startswith("t") and k[1:].isdigit():
                t = int(k[1:])
                q = f"'Parámetros'!$C${18 + t}"
                h = f"'Parámetros'!$D${18 + t}"
                et = (f'="Tramo {t}: "&IF({q}=0,"hasta "&{h}&" UIT",IF({h}="","más de "&{q}&" UIT",'
                      f'"más de "&{q}&" hasta "&{h}&" UIT"))&" ("&TEXT(\'Parámetros\'!$E${18 + t},"0%")&")"')
            bold = k in ("retc", "ia", "rb", "rn", "tot")
            estilo(ws.cell(row=f, column=2, value=et), F_BOLD if bold else F_BASE,
                   FILL_CLAVE if k == "retc" else None)
            for m in range(1, 13):
                fx = f'=IF($J$4="","",INDEX({CM}${C[k]}:${C[k]},$J$4+{m - 1}))'
                estilo(ws.cell(row=f, column=2 + m, value=fx), F_BOLD if bold else F_BASE,
                       FILL_CLAVE if k == "retc" else None, nf,
                       Alignment(horizontal="center", shrink_to_fit=True)
                       if nf in (None, NF_ENTERO, NF_PCT0) else None)
            if total:
                estilo(ws.cell(row=f, column=15, value=f'=IF($J$4="","",SUM(C{f}:N{f}))'), F_BOLD,
                       FILL_CLAVE if k == "retc" else FILL_AUX, nf)
            else:
                estilo(ws.cell(row=f, column=15), fill=FILL_AUX)
            self.ficha_filas[k] = f
            f += 1
        fres = self.ficha_filas["res"]
        rng = f"C{fres}:N{fres}"
        for texto, color, fondo in [("OK", "006100", "C6EFCE"), ("No corresponde", "006100", "C6EFCE"),
                                    ("Retuvo de menos", "9C0006", "FFC7CE"),
                                    ("Retuvo de más", "9C5700", "FFEB9C")]:
            ws.conditional_formatting.add(
                rng, FormulaRule(formula=[f'C{fres}="{texto}"'], font=Font(name=FONT, color=color, bold=True),
                                 fill=PatternFill("solid", fgColor=fondo)))
        ws.conditional_formatting.add(
            f"C{fh + 1}:N{f}", FormulaRule(formula=[f'C${self.ficha_filas["est"]}="No vinculado"'],
                                           font=Font(name=FONT, color="A6A6A6")))

        # Resumen anual por tramos
        f += 1
        ws.cell(row=f, column=2, value="IMPUESTO ANUAL POR TRAMOS (ingresos registrados del ejercicio)").font = F_SECCION
        f += 1
        RA = "'Resumen Anual'!"
        fila_res = f'$F$4+{self.t0 - 1}'

        def res(k):
            return f'=IF($F$4="","",INDEX({RA}${R[k]}:${R[k]},{fila_res}))'

        for j, h in enumerate(["Concepto", "Empleador (7 UIT)", "DJ anual (7 UIT + gastos)"]):
            encabezado(ws, f, 2 + j, h)
        ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=5)
        estilo(ws.cell(row=f, column=5), F_HDR, FILL_HDR)
        f += 1
        base = [("Renta bruta anual", res("rb"), res("rb")),
                ("(−) Deducción 7 UIT", res("d7"), res("d7")),
                ("(−) Deducción adicional por gastos (máx. 3 UIT)", 0, res("d3")),
                ("Renta neta imponible", res("rn"), res("rndj"))]
        f_rn = None
        for et, a, b in base:
            estilo(ws.cell(row=f, column=2, value=et), F_BOLD if et.startswith("Renta") else F_BASE)
            estilo(ws.cell(row=f, column=3, value=a), nf=NF_MONEDA)
            estilo(ws.cell(row=f, column=4, value=b), nf=NF_MONEDA)
            ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=5)
            if et == "Renta neta imponible":
                f_rn = f
            f += 1
        for t in range(1, 6):
            p = 18 + t
            d_, h_ = f"'Parámetros'!$F${p}", f"'Parámetros'!$G${p}"
            et = (f'="Tramo {t} ("&TEXT(\'Parámetros\'!$E${p},"0%")&"): "&IF({d_}=0,"hasta S/ "&TEXT({h_},"#,##0"),'
                  f'IF({h_}>=1E+14,"más de S/ "&TEXT({d_},"#,##0"),"S/ "&TEXT({d_},"#,##0")&" – "&TEXT({h_},"#,##0")))')
            estilo(ws.cell(row=f, column=2, value=et))
            for col, ref in ((3, f"C{f_rn}"), (4, f"D{f_rn}")):
                fx = (f'=IF({ref}="","",MAX(0,MIN(N({ref}),\'Parámetros\'!$G${p})-\'Parámetros\'!$F${p})'
                      f'*\'Parámetros\'!$E${p})')
                estilo(ws.cell(row=f, column=col, value=fx), nf=NF_MONEDA)
            ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=5)
            f += 1
        fin_tr = f - 1
        estilo(ws.cell(row=f, column=2, value="IMPUESTO ANUAL"), F_BOLD, FILL_CLAVE)
        estilo(ws.cell(row=f, column=3, value=f"=SUM(C{fin_tr - 4}:C{fin_tr})"), F_BOLD, FILL_CLAVE, NF_MONEDA)
        estilo(ws.cell(row=f, column=4, value=f"=SUM(D{fin_tr - 4}:D{fin_tr})"), F_BOLD, FILL_CLAVE, NF_MONEDA)
        ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=5)
        f_imp = f
        f += 1
        estilo(ws.cell(row=f, column=2, value="(−) Retenciones consideradas (empresa + otros)"))
        estilo(ws.cell(row=f, column=3, value=res("rcons")), nf=NF_MONEDA)
        estilo(ws.cell(row=f, column=4, value=res("rcons")), nf=NF_MONEDA)
        ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=5)
        f += 1
        estilo(ws.cell(row=f, column=2, value="Saldo: por retener / pagar (+) o en exceso / a favor (−)"),
               F_BOLD)
        estilo(ws.cell(row=f, column=3, value=f'=IF(C{f_imp}="","",C{f_imp}-C{f - 1})'), F_BOLD, nf=NF_MONEDA)
        estilo(ws.cell(row=f, column=4, value=f'=IF(D{f_imp}="","",D{f_imp}-D{f - 1})'), F_BOLD, nf=NF_MONEDA)
        ws.merge_cells(start_row=f, start_column=4, end_row=f, end_column=5)
        f += 1
        estilo(ws.cell(row=f, column=2, value="Estado de la validación"), F_BOLD)
        estilo(ws.cell(row=f, column=3, value=res("estado")), F_BOLD)
        ws.merge_cells(start_row=f, start_column=3, end_row=f, end_column=5)
        ws.freeze_panes = "C14"
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------------------------------------------------------- Instrucciones
    def hoja_instrucciones(self):
        ws = self.wb.active
        ws.title = "Instrucciones"
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 4
        ws.column_dimensions["C"].width = 120
        ws["B1"] = "VALIDADOR DE RETENCIONES DEL IMPUESTO A LA RENTA DE 5TA CATEGORÍA – PERÚ"
        ws["B1"].font = Font(name=FONT, size=16, bold=True, color=C_AZUL_OSC)
        ws["B2"] = ("Cálculo mensual (art. 40 del Reglamento de la LIR) y anual por tramos (art. 53 de la LIR), "
                    f"para varios trabajadores (capacidad: {self.N} trabajadores y {self.R} filas de ingresos).")
        ws["B2"].font = F_SUB
        contenido = [
            ("h", "¿QUÉ HACE ESTA PLANTILLA?"),
            ("•", "Calcula, para cada trabajador y cada mes, la retención de 5ta categoría según el art. 40 del "
                  "Reglamento de la LIR: proyección de la remuneración ordinaria, gratificaciones por percibir, "
                  "ingresos de meses anteriores, deducción de 7 UIT, escala progresiva por tramos y divisores."),
            ("•", "Aplica el procedimiento para ingresos extraordinarios o pagos particulares (utilidades, bonos, "
                  "reintegros, liquidaciones) en el mes en que se pagan."),
            ("•", "Acumula ingresos y retenciones de otros empleadores (empleo simultáneo o empleador anterior en "
                  "el mismo ejercicio)."),
            ("•", "Compara la retención calculada con la retención efectuada en planilla y marca las diferencias."),
            ("•", "Resume el impuesto anual por tramos y estima el saldo de la declaración jurada anual con la "
                  "deducción adicional de hasta 3 UIT."),
            ("h", "PASOS"),
            ("1", "Parámetros: indique el ejercicio y verifique la UIT (2026: S/ 5,500 – D.S. N° 301-2025-EF). "
                  "Para otro año, agregue su UIT en la tabla 'UIT por ejercicio'."),
            ("2", "Trabajadores: un trabajador por fila con código único, fechas de ingreso y cese, seguro de salud "
                  "(define el % de la bonificación extraordinaria) y opciones de proyección."),
            ("3", "Ingresos: una fila por concepto con código del trabajador, concepto (catálogo), origen, frecuencia, "
                  "monto por pago y meses. Las columnas Ene–Dic se llenan solas según la frecuencia:"),
            ("", "   – Mensual: todos los meses entre 'Mes inicio' y 'Mes fin' (vacío = diciembre)."),
            ("", "   – Bimestral / Trimestral / Cuatrimestral / Semestral: cada 2 / 3 / 4 / 6 meses contados desde "
                 "'Mes inicio', hasta 'Mes fin'."),
            ("", "   – Julio y Diciembre: gratificaciones ordinarias y su bonificación extraordinaria."),
            ("", "   – Pago único: solo en el 'Mes inicio' (utilidades, bonos, reintegros, liquidación, etc.)."),
            ("", "   – Ingresos variables (horas extras, comisiones): una fila 'Pago único' por mes. Aumento de "
                 "sueldo: cierre la fila con 'Mes fin' y agregue otra con el nuevo monto."),
            ("", "   – Registre el mes en que el ingreso se paga o se pone a disposición (art. 57 LIR: lo percibido)."),
            ("", "   – Otros empleadores: Origen = 'Otro empleador'. Sus retenciones (certificado de rentas y "
                 "retenciones) con el concepto 'Retención de 5ta efectuada por otro empleador'."),
            ("4", "Retenciones Planilla: digite la retención de 5ta descontada cada mes (boleta / PLAME). Deje en "
                  "blanco los meses sin información."),
            ("5", "Revise: 'Cálculo Mensual' (12 filas por trabajador, con filtro), 'Resumen Anual' (por trabajador "
                  "y por tramos) y 'Ficha Trabajador' (elija el código en la celda amarilla)."),
            ("h", "LEYENDA"),
            ("in", "Celda de ingreso de datos (fondo amarillo, texto azul)."),
            ("fx", "Fórmula: no modificar (texto negro; columnas auxiliares en gris)."),
            ("lk", "Vínculo a otra hoja (texto verde)."),
            ("ok", "Resultado clave (fondo verde)."),
            ("•", "Resultado de la validación: 'OK' (dentro de la tolerancia), 'Retuvo de menos', 'Retuvo de más', "
                  "'No corresponde' (no hay retención calculada ni registrada), 'Sin dato planilla' o "
                  "'Revisar: sin vínculo' (hay retención en un mes sin vínculo laboral)."),
            ("h", "PROCEDIMIENTO DE CÁLCULO MENSUAL (art. 40 del Reglamento de la LIR – D.S. N° 122-94-EF)"),
            ("a)", "Remuneración ordinaria del mes × meses que faltan para terminar el ejercicio (incluido el mes), "
                   "más las gratificaciones ordinarias (y bonificación extraordinaria) aún no percibidas."),
            ("b)", "Más las remuneraciones y demás ingresos afectos percibidos en los meses anteriores del ejercicio "
                   "(incluye otros empleadores)."),
            ("c)", "Menos 7 UIT = renta neta anual proyectada."),
            ("d)", "Impuesto anual proyectado con la escala progresiva acumulativa (art. 53 LIR): hasta 5 UIT 8%; "
                   "más de 5 hasta 20 UIT 14%; más de 20 hasta 35 UIT 17%; más de 35 hasta 45 UIT 20%; "
                   "más de 45 UIT 30%."),
            ("e)", "Retención del mes = (impuesto anual − retenciones previas) / divisor: enero a marzo ÷ 12 (sin "
                   "descontar); abril: − ret. ene–mar ÷ 9; mayo a julio: − ret. ene–abr ÷ 8; agosto: − ret. ene–jul "
                   "÷ 5; setiembre a noviembre: − ret. ene–ago ÷ 4; diciembre: regularización (impuesto anual − "
                   "ret. ene–nov)."),
            ("f)", "Ingresos extraordinarios del mes (utilidades, gratificaciones o bonificaciones extraordinarias, "
                   "reintegros, pagos particulares): retención adicional = impuesto con el extraordinario − "
                   "impuesto sin él, sobre la misma proyección. Se suma a la retención ordinaria del mes."),
            ("g)", "Mes de cese: impuesto sobre los ingresos percibidos hasta el cese, sin proyectar, menos todas las "
                   "retenciones anteriores. Si en diciembre o en el cese lo retenido supera el impuesto, se muestra "
                   "el exceso a devolver."),
            ("h", "OTROS EMPLEADORES"),
            ("•", "Empleo simultáneo: el trabajador informa al empleador que paga la mayor remuneración las rentas de "
                  "los demás, para que las acumule y efectúe la retención sobre el total. Por defecto, los ingresos "
                  "del otro empleador se suman conforme se perciben; si en Trabajadores marca 'Sí' en '¿Proyectar "
                  "rem. ordinaria de otro empleador?', la remuneración ordinaria del otro empleador también se "
                  "proyecta hasta diciembre."),
            ("•", "Empleador anterior en el mismo ejercicio: registre sus rentas y las retenciones del certificado de "
                  "rentas y retenciones; se suman a los ingresos de meses anteriores y se descuentan como "
                  "retenciones ya efectuadas."),
            ("h", "SUPUESTOS Y LIMITACIONES"),
            ("•", "Retenciones de meses anteriores: según el parámetro, se descuentan las de planilla (cuando se "
                  "digitaron) o las calculadas. Con 'Planilla (si existe)', un error de un mes se corrige en los "
                  "meses siguientes, como ocurre en la práctica."),
            ("•", "La deducción adicional de hasta 3 UIT (arrendamiento, honorarios médicos y profesionales, hoteles y "
                  "restaurantes, aportes a EsSalud de trabajadores del hogar, etc.) no la aplica el empleador: solo "
                  "se usa en el Resumen Anual para estimar el saldo de la DJ anual."),
            ("•", "Gratificación estimada: si la gratificación de julio o diciembre aún no se registró, se estima "
                  "con la remuneración computable del mes × (1 + % bonificación) × meses completos del semestre / 6, "
                  "solo en los meses anteriores al de pago. Registre la gratificación real cuando se pague."),
            ("•", "No incluye donaciones (art. 49 LIR), rentas de 4ta categoría, créditos ni convenios para evitar la "
                  "doble imposición."),
            ("•", "La retención calculada se redondea a 2 decimales; las diferencias dentro de la tolerancia "
                  "(Parámetros) se consideran OK."),
            ("•", "No elimine filas ni columnas: borre solo el contenido de las celdas amarillas. Para ampliar la "
                  "capacidad, vuelva a generar el archivo con generar_validador_5ta.py (--trabajadores, "
                  "--filas-ingresos)."),
            ("•", "Herramienta de apoyo: verifique los criterios con la normativa vigente y los pronunciamientos de "
                  "SUNAT."),
            ("h", "BASE LEGAL"),
            ("•", "TUO de la Ley del Impuesto a la Renta (D.S. N° 179-2004-EF): art. 18 inc. c) (inafectaciones), "
                  "art. 34 (rentas de 5ta categoría), art. 46 (deducción de 7 UIT y hasta 3 UIT adicionales), "
                  "art. 53 (escala), art. 57 (imputación por lo percibido), art. 75 (retención del empleador)."),
            ("•", "Reglamento de la LIR (D.S. N° 122-94-EF): art. 40 (procedimiento de retención) y normas sobre "
                  "trabajadores con más de un empleador y certificados de rentas y retenciones."),
            ("•", "D.S. N° 301-2025-EF: UIT 2026 = S/ 5,500."),
            ("•", "Ley N° 27735 (gratificaciones); Leyes N° 30334 y 30712 (bonificación extraordinaria)."),
        ]
        if self.ejemplos:
            contenido += [
                ("h", "DATOS DE EJEMPLO"),
                ("•", "Se incluyen 4 trabajadores de ejemplo (códigos EJ-001 a EJ-004) con sus ingresos y las "
                      "retenciones de planilla de enero a setiembre. Hay 2 diferencias intencionales (EJ-001 en "
                      "abril: no se retuvo sobre las utilidades; EJ-004 en marzo: se retuvo de más) para mostrar la "
                      "validación. Bórrelos (solo el contenido) antes de usar la plantilla."),
            ]
        f = 4
        estilos = {"in": (F_INPUT, FILL_INPUT, "Dato"), "fx": (F_BASE, FILL_AUX, "Fórmula"),
                   "lk": (F_LINK, None, "Vínculo"), "ok": (F_BOLD, FILL_CLAVE, "1,234.56")}
        for tipo, texto in contenido:
            if tipo == "h":
                f += 1
                c = ws.cell(row=f, column=2, value=texto)
                c.font = F_SECCION
                c.fill = FILL_SECCION
                ws.cell(row=f, column=3).fill = FILL_SECCION
            elif tipo in estilos:
                font, fill, muestra = estilos[tipo]
                c = ws.cell(row=f, column=2, value=muestra)
                estilo(c, font, fill, al=Alignment(horizontal="center"))
                ws.cell(row=f, column=3, value=texto).font = F_BASE
                ws.column_dimensions["B"].width = 9
            else:
                ws.cell(row=f, column=2, value=tipo).font = F_BOLD
                ws.cell(row=f, column=2).alignment = Alignment(horizontal="right", vertical="top")
                c = ws.cell(row=f, column=3, value=texto)
                c.font = F_BASE
                c.alignment = AL_WRAP
                lineas = max(1, -(-len(texto) // 125))
                ws.row_dimensions[f].height = 14 * lineas
            f += 1

    def generar(self, salida):
        self.hoja_instrucciones()
        self.hoja_parametros()
        self.hoja_conceptos()
        self.hoja_trabajadores()
        self.hoja_ingresos()
        self.hoja_retenciones()
        self.hoja_calculo()
        self.hoja_resumen()
        self.hoja_ficha()
        colores = {"Instrucciones": "404040", "Parámetros": "7F6000", "Conceptos": "7F6000",
                   "Trabajadores": "2E75B6", "Ingresos": "2E75B6", "Retenciones Planilla": "2E75B6",
                   "Cálculo Mensual": "548235", "Resumen Anual": "548235", "Ficha Trabajador": "548235"}
        for ws in self.wb.worksheets:
            ws.sheet_properties.tabColor = colores[ws.title]
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(salida)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--trabajadores", type=int, default=100)
    ap.add_argument("--filas-ingresos", type=int, default=1500)
    ap.add_argument("--anio", type=int, default=2026)
    ap.add_argument("--sin-ejemplos", action="store_true")
    ap.add_argument("--salida", default="Validador_Retenciones_5ta_Categoria.xlsx")
    a = ap.parse_args()
    Generador(a.trabajadores, a.filas_ingresos, not a.sin_ejemplos, a.anio).generar(a.salida)
    print(f"Archivo generado: {a.salida}")


if __name__ == "__main__":
    main()
